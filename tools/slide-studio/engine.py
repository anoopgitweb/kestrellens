"""Config-driven layouts shared by SVG preview and editable PowerPoint export."""
import json
import math
import textwrap
import base64
import charts
import images
import statistical
from pathlib import Path
from html import escape
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

ROOT = Path(__file__).parent
CATALOG = json.loads((ROOT / 'config/templates.json').read_text(encoding='utf-8'))
THEMES = json.loads((ROOT / 'config/themes.json').read_text(encoding='utf-8'))
TEMPLATES = {t['id']: t for t in CATALOG}

def validate(deck):
    if not isinstance(deck, dict) or deck.get('theme') not in THEMES:
        raise ValueError('Choose a valid theme.')
    slides = deck.get('slides')
    if not isinstance(slides, list) or not 1 <= len(slides) <= 60:
        raise ValueError('A presentation must have 1–60 slides.')
    for i, s in enumerate(slides):
        prefix = f'Slide {i + 1}: '
        if not isinstance(s, dict) or s.get('type') not in TEMPLATES:
            raise ValueError(prefix + 'unknown template.')
        for key, limit in [('title', 100), ('subtitle', 160), ('content', 20000 if s.get('chartType') in statistical.KINDS else 2400)]:
            value = s.get(key)
            if not isinstance(value, str) or len(value) > limit or any(ord(c) < 32 and c not in '\n\r\t' for c in value):
                raise ValueError(prefix + f'{key} must be text, at most {limit} characters.')
        if not s['title'].strip():
            raise ValueError(prefix + 'add a title.')
        t = TEMPLATES[s['type']]
        rows = parse(s)
        if t['layout'] == 'chart':
            charts.data(s)
        elif t['fields']:
            maximum = 8 if t['layout'] in ('chart', 'table', 'summary') else 6
            if not 1 <= len(rows) <= maximum:
                raise ValueError(prefix + f'enter 1–{maximum} rows.')
            for row in rows:
                if len(row) != len(t['fields']) or any(not cell for cell in row):
                    raise ValueError(prefix + 'each row needs: ' + ' | '.join(t['fields']))
                if any(len(cell) > (180 if t['layout'] == 'summary' else 100) for cell in row):
                    raise ValueError(prefix + 'shorten each cell to fit the slide.')
            if t['layout'] == 'chart':
                try:
                    values = [float(r[1]) for r in rows]
                    if any(not math.isfinite(v) or abs(v) > 1e12 for v in values):
                        raise ValueError()
                except ValueError:
                    raise ValueError(prefix + 'chart values must be finite numbers between -1e12 and 1e12.')
        elif len(s['content']) > 250:
            raise ValueError(prefix + 'closing text must be at most 250 characters.')
        scene(s, THEMES[deck['theme']], i + 1)  # Fail before export on text overflow.
    return deck

def parse(s):
    return [[cell.strip() for cell in line.split('|')] for line in s['content'].splitlines() if line.strip()]

def scene(s, theme, number):
    objects = []
    def box(x, y, w, h, color):
        objects.append(dict(kind='rect', x=x, y=y, w=w, h=h, color=theme.get(color, color)))
    def text(value, x, y, w, h, size=22, color='text', bold=False):
        # Explicit wrapping and fixed line heights are shared by both renderers.
        for candidate in range(size, 13, -1):
            lines = []
            for paragraph in str(value).splitlines() or ['']:
                lines.extend(textwrap.wrap(paragraph, max(1, int(w / (candidate * .57))), break_long_words=True) or [''])
            if len(lines) * candidate * 1.28 <= h:
                break
        else:
            raise ValueError('Text exceeds the slide layout. Shorten the content or split it across slides.')
        objects.append(dict(kind='text', x=x, y=y, w=w, h=h, lines=lines, size=candidate, color=theme[color], bold=bold))
    box(0, 0, 1200, 675, 'background')
    t = TEMPLATES[s['type']]
    layout = t['layout']
    if layout == 'hero':
        # Reference-inspired cover: generous title scale, quiet metadata, no card chrome.
        if theme['background'] != 'FFFFFF':
            box(970, 0, 230, 675, 'panel')
        text('STATISTICAL ANALYSIS REPORT', 78, 72, 650, 36, 16, 'accent', True)
        text(s['title'], 78, 166, 1030, 150, 52, bold=True)
        text(s['subtitle'], 80, 350, 990, 70, 26, 'muted')
        text(s['content'], 80, 492, 980, 70, 19, 'muted')
    else:
        section = {'summary':'OVERVIEW','cards':'KEY METRICS','comparison':'COMPARISON','chart':'STATISTICAL ANALYSIS','process':'PROCESS','timeline':'ROADMAP','table':'ACTION PLAN'}.get(layout,'ANALYSIS')
        text(section, 60, 31, 700, 30, 15, 'accent', True)
        text(s['title'], 60, 70, 1000, 78, 42, bold=True)
        text(s['subtitle'], 60, 135, 980, 48, 18, 'muted')
        text(f'{number:02} / 10', 1050, 34, 100, 28, 15, 'muted')
        rows = parse(s)
        n = len(rows)
        if layout == 'summary':
            h = 390 / n
            for i, row in enumerate(rows):
                text(f'{i+1:02}', 60, 220+i*h, 55, h-6, 24, 'accent', True)
                text(row[0], 140, 220+i*h, 990, h-6, 26)
        elif layout == 'timeline':
            h = 380/n
            box(80, 230, 3, 380-h+20, 'accent')
            for i, row in enumerate(rows):
                y = 225+i*h
                box(73, y+8, 17, 17, 'accent')
                text(row[0], 115, y, 200, h-10, 25, 'accent', True)
                text(row[1], 345, y, 780, h-10, 25)
        elif layout == 'process':
            h = 380/n
            for i, row in enumerate(rows):
                y = 220+i*h
                text(f'{i+1:02}', 65, y, 65, h-10, 25, 'accent', True)
                text(row[0], 165, y, 300, h-10, 25, 'text', True)
                text(row[1], 500, y, 620, h-10, 24)
                if i < n-1:
                    box(90, y+32, 2, max(2,h-38), 'accent')
        elif layout in ('cards', 'comparison'):
            cols = min(n, 3)
            count_rows = math.ceil(n/cols)
            w, h = 1080/cols, 385/count_rows
            for i, row in enumerate(rows):
                x, y = 60+(i%cols)*w, 220+(i//cols)*h
                box(x, y, w-18, h-18, 'panel')
                headline = row[0] if layout != 'process' else f'{i+1}. {row[0]}'
                text(headline, x+20, y+18, w-58, (h-48)*.44, 36 if layout == 'cards' else 26, 'accent', True)
                text('\n'.join(row[1:]), x+20, y+20+(h-48)*.44, w-58, (h-48)*.56, 23)
        elif layout in ('timeline', 'table'):
            widths = [540, 280, 260] if layout == 'table' else [220, 860]
            data = [t['fields']] + rows
            h = 390/len(data)
            for i, row in enumerate(data):
                box(60, 215+i*h, 1080, h-3, 'panel' if i%2 == 0 else 'background')
                x = 60
                for cell, w in zip(row, widths):
                    text(cell, x+14, 220+i*h, w-28, h-12, 22, 'accent' if i == 0 else 'text', i == 0)
                    x += w
        elif layout == 'chart':
            if s.get('chartType') in statistical.KINDS:
                objects.extend(statistical.objects(s,theme))
            else:
                objects.append(dict(kind='chart', x=60,y=210,w=1080,h=410,source=s,theme=theme))
    if s.get('image'):
        obj,reserved=images.object_for(s)
        scale=(1080-reserved-30)/1080
        offset=reserved+30 if s.get('imagePosition','right')=='left' else 0
        for o in objects[1:]:
            o['x']=60+(o['x']-60)*scale+offset
            o['w']*=scale
            if o['kind']=='text':
                o['size']*=scale
        objects.append(obj)
    text(f'{number:02}', 1090, 632, 60, 26, 16, 'muted')
    return objects

def svg(objects):
    parts = ['<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1200 675" role="img" aria-label="Slide preview">']
    for o in objects:
        if o['kind'] == 'rect':
            parts.append(f'<rect x="{o["x"]}" y="{o["y"]}" width="{o["w"]}" height="{o["h"]}" fill="#{o["color"]}"/>')
        elif o['kind']=='line':
            parts.append(f'<line x1="{o["x"]}" y1="{o["y"]}" x2="{o["x"]+o["w"]}" y2="{o["y"]+o["h"]}" stroke="#{o["color"]}" stroke-width="{o["width"]}"/>')
        elif o['kind']=='ellipse':
            parts.append(f'<ellipse cx="{o["x"]+o["w"]/2}" cy="{o["y"]+o["h"]/2}" rx="{o["w"]/2}" ry="{o["h"]/2}" fill="#{o["color"]}"/>')
        elif o['kind'] in ('chart','image'):
            raw=charts.preview(o['source'],o['theme']).encode() if o['kind']=='chart' else o['raw']
            mime='image/svg+xml' if o['kind']=='chart' else 'image/png'
            uri='data:'+mime+';base64,'+base64.b64encode(raw).decode()
            parts.append(f'<image x="{o["x"]}" y="{o["y"]}" width="{o["w"]}" height="{o["h"]}" preserveAspectRatio="none" href="{uri}"/>')
        else:
            for i, line in enumerate(o['lines']):
                parts.append(f'<text x="{o["x"]}" y="{o["y"]+o["size"]+i*o["size"]*1.28}" font-family="Arial, sans-serif" font-size="{o["size"]}" font-weight="{700 if o["bold"] else 400}" fill="#{o["color"]}">{escape(line)}</text>')
    return ''.join(parts)+'</svg>'

def powerpoint(deck):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    unit = lambda v: Inches(v/90)
    for i, s in enumerate(deck['slides']):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for o in scene(s, THEMES[deck['theme']], i+1):
            if o['kind'] in ('rect','ellipse'):
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE if o['kind']=='rect' else MSO_SHAPE.OVAL, unit(o['x']), unit(o['y']), unit(o['w']), unit(o['h']))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(o['color'])
                shape.line.fill.background()
            elif o['kind']=='line':
                shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,unit(o['x']),unit(o['y']),unit(o['x']+o['w']),unit(o['y']+o['h']))
                shape.line.color.rgb=RGBColor.from_string(o['color'])
                shape.line.width=Pt(o['width']*.8)
            elif o['kind']=='chart':
                charts.add_chart(slide,o,unit)
            elif o['kind']=='image':
                slide.shapes.add_picture(BytesIO(o['raw']),unit(o['x']),unit(o['y']),unit(o['w']),unit(o['h']))
            else:
                shape = slide.shapes.add_textbox(unit(o['x']), unit(o['y']), unit(o['w']), unit(o['h']))
                tf = shape.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False
                for j, line in enumerate(o['lines']):
                    p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                    p.text = line
                    p.font.name = 'Arial'
                    p.font.size = Pt(o['size']*.8)
                    p.font.bold = o['bold']
                    p.font.color.rgb = RGBColor.from_string(o['color'])
                    p.space_before = p.space_after = Pt(0)
                    p.line_spacing = Pt(o['size']*.8*1.28)
    output = BytesIO()
    prs.save(output)
    return output.getvalue()
