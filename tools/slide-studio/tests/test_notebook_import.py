import sys
import zipfile
import base64
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '.packages'))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import unittest
from pptx import Presentation
from app import app
from engine import scene, THEMES


def workbook():
    headers = ['Notebook', 'Chapter', 'Page Heading', 'Page Details', 'Video URL',
               'Local Video URL', 'Diagram', 'Page Order']
    rows = [
        ['Imported Deck', 'Foundations', 'Second page', 'Second body', '', '', '', '2'],
        ['Imported Deck', 'Foundations', 'First page', 'First body', 'https://example.com/video',
         'lesson.mp4', '', '1'],
    ]
    def cell(column, row, value):
        letters = chr(65 + column)
        return f'<c r="{letters}{row}" t="inlineStr"><is><t>{value}</t></is></c>'
    sheet_rows = []
    for number, values in enumerate([headers, *rows], 1):
        sheet_rows.append(f'<row r="{number}">' + ''.join(cell(i, number, value) for i, value in enumerate(values)) + '</row>')
    sheet = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
             'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
             '<sheetData>' + ''.join(sheet_rows) + '</sheetData><drawing r:id="rId1"/></worksheet>')
    output = BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
        archive.writestr('_rels/.rels', '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        archive.writestr('xl/workbook.xml', '<?xml version="1.0"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Notebook" sheetId="1" r:id="rId1"/></sheets></workbook>')
        archive.writestr('xl/_rels/workbook.xml.rels', '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
        archive.writestr('xl/worksheets/sheet1.xml', sheet)
        archive.writestr('xl/worksheets/_rels/sheet1.xml.rels', '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>')
        archive.writestr('xl/drawings/drawing1.xml', '<?xml version="1.0"?><xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><xdr:oneCellAnchor><xdr:from><xdr:col>6</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>2</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:ext cx="952500" cy="952500"/><xdr:pic><xdr:blipFill><a:blip r:embed="rId1"/></xdr:blipFill></xdr:pic><xdr:clientData/></xdr:oneCellAnchor></xdr:wsDr>')
        archive.writestr('xl/drawings/_rels/drawing1.xml.rels', '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/></Relationships>')
        archive.writestr('xl/media/image1.png', base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII='))
    return output.getvalue()


class NotebookImportTests(unittest.TestCase):
    def setUp(self):
        self.client = app.test_client()

    def test_import_orders_pages_and_generates_powerpoint(self):
        response = self.client.post('/api/import-notebook', data={
            'file': (BytesIO(workbook()), 'notebook.xlsx')
        }, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 200, response.data)
        deck = response.get_json()
        self.assertEqual(deck['name'], 'Imported Deck')
        self.assertEqual([slide['title'] for slide in deck['slides']],
                         ['Imported Deck', 'Foundations', 'First page', 'Second page'])
        self.assertEqual(deck['slides'][0]['eyebrow'], 'NOTEBOOK PRESENTATION')
        self.assertEqual(deck['slides'][1]['eyebrow'], 'CHAPTER')
        self.assertEqual(deck['slides'][2]['sourceUrl'], 'https://example.com/video')
        self.assertEqual(deck['slides'][2]['localVideoReference'], 'lesson.mp4')
        self.assertTrue(deck['slides'][2]['image'].startswith('data:image/png;base64,'))
        objects = scene(deck['slides'][2], THEMES[deck['theme']], 3)
        imported_image = next(item for item in objects if item['kind'] == 'image')
        self.assertEqual(imported_image['y'], 225)
        self.assertTrue(any('LOOK · CONNECT · UNDERSTAND' in item.get('lines', []) for item in objects))
        generated = self.client.post('/api/generate', json=deck)
        self.assertEqual(generated.status_code, 200, generated.data[:400])
        presentation = Presentation(BytesIO(generated.data))
        self.assertEqual(len(presentation.slides), 5)
        self.assertTrue(any(shape.shape_type == 13 for shape in presentation.slides[2].shapes))
        source_picture = next(shape for shape in presentation.slides[2].shapes if shape.shape_type == 13)
        detail_slide = presentation.slides[4]
        self.assertEqual(detail_slide._element.get('show'), '0')
        self.assertEqual(source_picture.click_action.target_slide.slide_id, detail_slide.slide_id)
        detail_picture = next(shape for shape in detail_slide.shapes if shape.shape_type == 13)
        self.assertEqual(detail_picture.click_action.target_slide.slide_id, presentation.slides[2].slide_id)
        link_runs = [run for shape in presentation.slides[2].shapes if shape.has_text_frame
                     for paragraph in shape.text_frame.paragraphs for run in paragraph.runs
                     if run.hyperlink.address]
        self.assertEqual(link_runs[0].hyperlink.address, 'https://example.com/video')

    def test_rejects_wrong_file_type(self):
        response = self.client.post('/api/import-notebook', data={
            'file': (BytesIO(b'not excel'), 'notebook.csv')
        }, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 400)


if __name__ == '__main__':
    unittest.main()
