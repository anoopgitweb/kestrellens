import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '.packages'))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import unittest
from io import BytesIO
from pptx import Presentation
from app import app
from engine import CATALOG, THEMES

def sample():
    return {'theme':'ocean','slides':[{'type':t['id'], **{k:t[k] for k in ('title','subtitle','content')}} for t in CATALOG]}

class AppTests(unittest.TestCase):
    def setUp(self): self.client = app.test_client()
    def test_all_templates_and_themes_export(self):
        for theme in THEMES:
            deck=sample(); deck['theme']=theme
            r=self.client.post('/api/preview',json=deck)
            self.assertEqual(r.status_code,200,r.data)
            self.assertEqual(len(r.json['slides']),9)
            r=self.client.post('/api/generate',json=deck)
            self.assertEqual(r.status_code,200,r.data[:300])
            p=Presentation(BytesIO(r.data))
            self.assertEqual(len(p.slides),9)
            for slide,source in zip(p.slides,deck['slides']):
                text=' '.join(s.text for s in slide.shapes if s.has_text_frame)
                self.assertIn(source['title'],text)
    def test_bad_input(self):
        for bad in [None,{}, {'theme':'unknown','slides':[]}, {'theme':'ocean','slides':[]}]:
            self.assertEqual(self.client.post('/api/generate',json=bad).status_code in (400,415),True)
        deck=sample();deck['slides'][4]['content']='Jan | NaN'
        self.assertEqual(self.client.post('/api/generate',json=deck).status_code,400)
        deck=sample();deck['slides'][0]['title']='x'*101
        self.assertEqual(self.client.post('/api/generate',json=deck).status_code,400)
    def test_negative_zero_chart_and_escape(self):
        deck=sample();deck['slides'][4]['content']='A | -20\nB | 0\nC | 10'
        deck['slides'][0]['title']='<script>alert(1)</script>'
        r=self.client.post('/api/preview',json=deck)
        self.assertEqual(r.status_code,200)
        self.assertNotIn('<script>',r.json['slides'][0])
        self.assertIn('&lt;script&gt;',r.json['slides'][0])
        self.assertEqual(self.client.post('/api/generate',json=deck).status_code,200)
    def test_order_and_delete(self):
        deck=sample();deck['slides']=[deck['slides'][8],deck['slides'][0]]
        p=Presentation(BytesIO(self.client.post('/api/generate',json=deck).data))
        self.assertEqual(len(p.slides),2)
        self.assertIn('Thank you',' '.join(s.text for s in p.slides[0].shapes if s.has_text_frame))
if __name__=='__main__': unittest.main()
