"""Offline chart validation, SVG preview, and native editable Office charts."""
import math
import statistical
from html import escape
from copy import deepcopy
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement

KINDS = {'column':'Column', 'bar':'Horizontal bar', 'line':'Line', 'area':'Area',
         'stacked':'Stacked column', 'pie':'Pie', 'donut':'Donut', 'scatter':'Scatter', 'combo':'Column + line'}
COLORS = ['2563EB','10B981','F59E0B','E879F9']

def data(s):
    kind = s.get('chartType','column')
    if kind in statistical.KINDS:
        return statistical.analyze(s)
    if kind not in KINDS: raise ValueError('Choose a supported chart type.')
    rows = [[c.strip() for c in r.split('|')] for r in s['content'].splitlines() if r.strip()]
    if not 1 <= len(rows) <= 12: raise ValueError('Charts need 1–12 data rows.')
    width = len(rows[0])
    if not 2 <= width <= 5 or any(len(r)!=width or any(not c for c in r) for r in rows):
        raise ValueError('Use Category | Value 1 | Value 2, with the same number of values on every row (up to 4 series).')
    if any(len(r[0]) > 24 for r in rows): raise ValueError('Chart category labels must be at most 24 characters.')
    try:
        vals = [[float(v) for v in r[1:]] for r in rows]
        if any(not math.isfinite(v) or abs(v)>1e12 for r in vals for v in r): raise ValueError()
        if kind=='scatter' and any(not math.isfinite(float(r[0])) or abs(float(r[0]))>1e12 for r in rows): raise ValueError()
    except ValueError: raise ValueError('Chart values must be finite numbers between -1e12 and 1e12. Scatter also needs numeric X values.')
    if kind in ('pie','donut') and (width!=2 or any(r[0]<0 for r in vals) or sum(r[0] for r in vals)<=0):
        raise ValueError('Pie and donut charts need one nonnegative series with a positive total.')
    if kind=='combo' and width!=3: raise ValueError('Combination charts need exactly two series: Category | Column value | Line value.')
    names = s.get('seriesNames','')
    if not isinstance(names,str) or len(names)>160: raise ValueError('Series names must be text of at most 160 characters.')
    names = [n.strip() for n in names.split('|')] if names.strip() else [f'Series {i+1}' for i in range(width-1)]
    if len(names)!=width-1 or any(not n or len(n)>32 for n in names): raise ValueError('Enter one series name per value column, separated by | (32 characters each).')
    return kind, [r[0] for r in rows], vals, names

def preview(s, theme):
    kind, labels, rows, names = data(s)
    colors = [theme['accent']]+COLORS[1:]
    out=[]
    def txt(x,y,v,size=15): out.append(f'<text x="{x}" y="{y}" fill="#{theme["text"]}" font-family="Arial" font-size="{size}">{escape(str(v))}</text>')
    def line(x,y,x2,y2,color,width=2): out.append(f'<line x1="{x}" y1="{y}" x2="{x2}" y2="{y2}" stroke="#{color}" stroke-width="{width}"/>')
    def rect(x,y,w,h,color): out.append(f'<rect x="{x}" y="{y}" width="{max(0,w)}" height="{max(0,h)}" fill="#{color}"/>')
    if kind in ('pie','donut'):
        total=sum(r[0] for r in rows);angle=-math.pi/2
        palette=colors+['F97316','06B6D4','8B5CF6','EF4444','84CC16','EC4899','14B8A6','A16207']
        for i,(label,row) in enumerate(zip(labels,rows)):
            a2=angle+2*math.pi*row[0]/total
            if row[0]==total:
                out.append(f'<circle cx="340" cy="185" r="150" fill="#{palette[i]}"/>')
            elif row[0]>0:
                x,y=340+150*math.cos(angle),185+150*math.sin(angle)
                x2,y2=340+150*math.cos(a2),185+150*math.sin(a2)
                out.append(f'<path d="M340 185 L{x} {y} A150 150 0 {int(a2-angle>math.pi)} 1 {x2} {y2} Z" fill="#{palette[i]}"/>')
            angle=a2
            rect(565,12+i*29,13,13,palette[i]);txt(590,24+i*29,f'{label}: {row[0]:g} ({row[0]/total:.0%})')
        if kind=='donut': out.append(f'<circle cx="340" cy="185" r="80" fill="#{theme["background"]}"/>')
    else:
        vals=[v for r in rows for v in r]
        if kind=='stacked': vals=[sum(max(0,v) for v in r) for r in rows]+[sum(min(0,v) for v in r) for r in rows]
        lo,hi=min(0,min(vals)),max(0,max(vals));hi=hi if hi!=lo else lo+1
        py=lambda v: 315-(v-lo)/(hi-lo)*265
        n=len(rows);step=880/n
        for j in range(5):
            v=lo+(hi-lo)*j/4;y=py(v);line(85,y,990,y,theme['panel'],1);txt(0,y+5,f'{v:.3g}')
        zero=py(0)
        if kind=='bar':
            out=[];px=lambda v:190+(v-lo)/(hi-lo)*740
            for j in range(5):
                v=lo+(hi-lo)*j/4;x=px(v);line(x,20,x,315,theme['panel'],1);txt(x-15,345,f'{v:.3g}')
            for i,r in enumerate(rows):
                y=25+i*285/n;txt(0,y+15,labels[i]);h=240/n/len(names)
                for k,v in enumerate(r):rect(min(px(0),px(v)),y+k*h,abs(px(v)-px(0)),h-2,colors[k])
        else:
            xs=[85+step*(i+.5) for i in range(n)]
            if kind=='scatter':
                xv=list(map(float,labels));xmin,xmax=min(xv),max(xv)
                if xmin==xmax:xmin-=1;xmax+=1
                xs=[100+(v-xmin)/(xmax-xmin)*860 for v in xv]
                for j in range(5):txt(85+j*220,346,f'{xmin+(xmax-xmin)*j/4:.3g}')
            else:
                for i,label in enumerate(labels):
                    for j in range(0,len(label),12):txt(xs[i]-25,340+j//12*17,label[j:j+12],13)
            for k,name in enumerate(names):
                if kind in ('line','area','scatter') or kind=='combo' and k==1:
                    pts=' '.join(f'{x},{py(r[k])}' for x,r in zip(xs,rows))
                    if kind=='area':out.append(f'<polygon points="{xs[0]},{zero} {pts} {xs[-1]},{zero}" fill="#{colors[k]}" opacity=".22"/>')
                    if kind!='scatter':out.append(f'<polyline points="{pts}" fill="none" stroke="#{colors[k]}" stroke-width="3"/>')
                    for x,r in zip(xs,rows):out.append(f'<circle cx="{x}" cy="{py(r[k])}" r="4" fill="#{colors[k]}"/>')
                else:
                    for i,r in enumerate(rows):
                        v=r[k];base=sum(vv for vv in r[:k] if (vv>=0)==(v>=0)) if kind=='stacked' else 0
                        width=step*.65/(1 if kind in ('stacked','combo') else len(names))
                        x=xs[i]-step*.325+(0 if kind in ('stacked','combo') else k*width)
                        rect(x,min(py(base),py(base+v)),width-2,abs(py(base+v)-py(base)),colors[k])
        for k,name in enumerate(names):rect(70+k*240,380,12,12,colors[k]);txt(90+k*240,392,name)
    return '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1080 410">'+''.join(out)+'</svg>'

def add_chart(slide,o,unit):
    kind,labels,rows,names=data(o['source'])
    cd=XyChartData() if kind=='scatter' else CategoryChartData()
    if kind=='scatter':
        for k,name in enumerate(names):
            series=cd.add_series(name)
            for label,row in zip(labels,rows):series.add_data_point(float(label),row[k])
    else:
        cd.categories=labels
        for k,name in enumerate(names):cd.add_series(name,[r[k] for r in rows])
    types={'column':XL_CHART_TYPE.COLUMN_CLUSTERED,'bar':XL_CHART_TYPE.BAR_CLUSTERED,'line':XL_CHART_TYPE.LINE_MARKERS,'area':XL_CHART_TYPE.AREA,'stacked':XL_CHART_TYPE.COLUMN_STACKED,'pie':XL_CHART_TYPE.PIE,'donut':XL_CHART_TYPE.DOUGHNUT,'scatter':XL_CHART_TYPE.XY_SCATTER,'combo':XL_CHART_TYPE.COLUMN_CLUSTERED}
    chart=slide.shapes.add_chart(types[kind],unit(o['x']),unit(o['y']),unit(o['w']),unit(o['h']),cd).chart
    if kind=='combo':
        plot=chart._chartSpace.xpath('.//c:barChart')[0]
        second=plot.xpath('./c:ser')[1];plot.remove(second)
        lp=OxmlElement('c:lineChart');group=OxmlElement('c:grouping');group.set('val','standard');lp.append(group);lp.append(second)
        for axis in plot.xpath('./c:axId'):lp.append(deepcopy(axis))
        plot.addnext(lp)
    chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout=False
    chart.font.name='Arial';chart.font.size=Pt(12);chart.font.color.rgb=RGBColor.from_string(o['theme']['text'])
    colors=[o['theme']['accent']]+COLORS[1:]
    for i,series in enumerate(chart.series):
        series.format.fill.solid();series.format.fill.fore_color.rgb=RGBColor.from_string(colors[i%4])
        series.format.line.color.rgb=RGBColor.from_string(colors[i%4])
    if kind not in ('pie','donut'):
        for axis in (chart.category_axis,chart.value_axis):axis.tick_labels.font.color.rgb=RGBColor.from_string(o['theme']['text'])
    else:
        chart.plots[0].vary_by_categories=True
        chart.plots[0].has_data_labels=True
        chart.plots[0].data_labels.show_percentage=True
